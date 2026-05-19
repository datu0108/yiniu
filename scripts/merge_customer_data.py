#!/usr/bin/env python3
"""亿牛广告看板：按今日推广数据关联客户与运营信息，输出合并报表。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 主管姓名 -> 战区
SUPERVISOR_WARZONE = {
    "褚鹏松": "柯虞战区",
    "余民": "新嵊战区",
    "林峰": "第一战区",
    "卢军": "第一战区",
}

# 数据文件路径（可按需修改）
DATA_DIR = Path("/Users/datu/PyCharmMiscProject/yiniu文件")
TABLE1_PATH = DATA_DIR / "5.3.xlsx"
TABLE2_PATH = DATA_DIR / "服务中客户(1).xlsx"
TABLE3_PATH = DATA_DIR / "对应运营1.xlsx"
OUTPUT_PATH = DATA_DIR / "合并结果.xlsx"
SHEET_BASE = "基础信息表"
SHEET_STATS = "统计数据"
SHEET_SALES_REPORT = "销售汇报"
SHEET_OPS_REPORT = "运营汇报"

# 表1 统计用字段
COL_CASH_DAY = "当天P现金消耗_元_"
COL_BUDGET = "当日整体p预算_元"
COL_FIN_DAY = "当天P财务消耗_元_"
COL_CASH_MONTH = "当月P现金消耗_元_"
COL_BALANCE = "账户现金余额_元"
COL_PREMIUM_PRODUCTS = "推广中实力优品数"
COL_SHOP_PRODUCTS = "店铺产品数"
COL_SHOP_PREMIUM = "店铺实力优品数"
COL_INQUIRY = "总询盘数_CPF分母"
ACTIVE_THRESHOLD = 800
ACTIVE_CUSTOMER_BASE = 650  # 活跃客户率分母
BALANCE_DAYS_THRESHOLD = 14
PREMIUM_PRODUCTS_MIN = 30
OPPORTUNITY_COST_THRESHOLD = 500
COL_SALES = "销售姓名"
WARZONE_ORDER = ["柯虞战区", "新嵊战区", "第一战区"]
COL_OPS = "运营"
COL_OPS_REGION = "运营区域"
OPS_REGION_ORDER = ["柯桥", "嵊州"]
# 运营 -> 区域（柯桥 / 嵊州）
OPS_TO_REGION = {
    "何晓宇": "柯桥",
    "张夏楠": "柯桥",
    "陶一晴": "柯桥",
    "赵毅": "柯桥",
    "徐佳铭": "柯桥",
    "宣雪颖": "柯桥",
    "张甜怡": "嵊州",
    "周杭露": "嵊州",
    "吴梦婷": "嵊州",
    "徐姗姗": "嵊州",
}
REPORT_METRICS = [
    "总消耗",
    "消耗使用率",
    "活跃客户数",
    "余额不足客户数",
    "商机成本高于500客户数",
]


def normalize_id(series: pd.Series) -> pd.Series:
    """统一 ID 格式，避免空格/大小写导致匹配失败。"""
    return series.astype(str).str.strip().str.lower()


def assign_warzone(supervisor: pd.Series) -> pd.Series:
    """根据主管姓名划分战区。"""
    name = supervisor.astype(str).str.strip()
    return name.map(SUPERVISOR_WARZONE)


def assign_ops_region(ops: pd.Series) -> pd.Series:
    """根据运营姓名划分柯桥 / 嵊州。"""
    name = ops.astype(str).str.strip()
    return name.map(OPS_TO_REGION)


def to_num(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0)


def set_column_widths(worksheet, df: pd.DataFrame, min_width: float = 16, max_width: float = 50) -> None:
    """按列内容自动加宽，中文列适当放宽。"""
    for idx, col in enumerate(df.columns, 1):
        header_len = len(str(col))
        values = df[col].fillna("").astype(str)
        content_len = int(values.str.len().max()) if len(values) else 0
        width = min(max(max(header_len, content_len) * 1.3 + 2, min_width), max_width)
        if col == "备注":
            width = min(max(width, 36), 60)
        worksheet.column_dimensions[get_column_letter(idx)].width = width


def prepare_row_metrics(df1: pd.DataFrame) -> pd.DataFrame:
    """计算行级指标及是否符合统计筛选条件。"""
    cash_day = to_num(df1[COL_CASH_DAY])
    fin_day = to_num(df1[COL_FIN_DAY])
    cash_month = pd.to_numeric(df1[COL_CASH_MONTH], errors="coerce")
    balance = to_num(df1[COL_BALANCE])
    premium_products = to_num(df1[COL_PREMIUM_PRODUCTS])
    inquiry = to_num(df1[COL_INQUIRY])

    coupon_day = fin_day - cash_day
    balance_days = balance / cash_day.replace(0, float("nan"))
    opportunity_cost = cash_day / inquiry.replace(0, float("nan"))

    is_active = cash_month >= ACTIVE_THRESHOLD
    is_low_balance = (cash_day > 0) & (balance_days <= BALANCE_DAYS_THRESHOLD)
    is_premium_insufficient = (cash_day > 0) & (premium_products < PREMIUM_PRODUCTS_MIN)
    is_high_opportunity_cost = (inquiry > 0) & (opportunity_cost > OPPORTUNITY_COST_THRESHOLD)

    tag_df = pd.DataFrame(
        {
            "活跃客户": is_active,
            "余额不足14天": is_low_balance,
            "优品加推不足": is_premium_insufficient,
            "商机成本高于500元": is_high_opportunity_cost,
        }
    )
    tags = tag_df.apply(
        lambda row: "；".join(tag_df.columns[row.values]) if row.any() else "",
        axis=1,
    )

    return pd.DataFrame(
        {
            COL_BALANCE: balance.round(2),
            "当日卡券消耗": coupon_day.round(2),
            "余额可用天数": balance_days.round(2),
            "当日商机成本": opportunity_cost.round(2),
            COL_PREMIUM_PRODUCTS: premium_products,
            "备注": tags,
            "_is_active": is_active,
            "_is_low_balance": is_low_balance,
            "_is_premium_insufficient": is_premium_insufficient,
            "_is_high_opportunity_cost": is_high_opportunity_cost,
        },
        index=df1.index,
    )


def _calc_metric_rows(df1: pd.DataFrame) -> list[tuple[str, str, float]]:
    """计算单日推广数据的统计指标。"""
    cash_day = to_num(df1[COL_CASH_DAY])
    budget_day = to_num(df1[COL_BUDGET])
    fin_day = to_num(df1[COL_FIN_DAY])
    cash_month = pd.to_numeric(df1[COL_CASH_MONTH], errors="coerce")
    premium_products = to_num(df1[COL_PREMIUM_PRODUCTS])
    inquiry = to_num(df1[COL_INQUIRY])
    row_metrics = prepare_row_metrics(df1)

    total_consume = cash_day.sum()
    total_budget = budget_day.sum()
    usage_rate = total_consume / total_budget if total_budget else 0
    coupon_consume = (fin_day - cash_day).sum()
    active_count = int((cash_month >= ACTIVE_THRESHOLD).sum())
    active_rate = active_count / ACTIVE_CUSTOMER_BASE

    return [
        ("总消耗", f"{COL_CASH_DAY} 求和", round(total_consume, 2)),
        ("总预算", f"{COL_BUDGET} 求和", round(total_budget, 2)),
        ("消耗使用率", "总消耗 / 总预算", round(usage_rate, 4)),
        ("当日卡券消耗", f"{COL_FIN_DAY} - {COL_CASH_DAY} 求和", round(coupon_consume, 2)),
        ("活跃客户数", f"{COL_CASH_MONTH} >= {ACTIVE_THRESHOLD} 的客户数", float(active_count)),
        ("活跃客户率", f"活跃客户数 / {ACTIVE_CUSTOMER_BASE}", round(active_rate, 4)),
        (
            "余额不足14天客户数",
            f"{COL_BALANCE} / {COL_CASH_DAY} <= {BALANCE_DAYS_THRESHOLD}（且当日消耗>0）",
            float(int(row_metrics["_is_low_balance"].sum())),
        ),
        ("直通车优品加推数", f"{COL_PREMIUM_PRODUCTS} 求和", float(int(premium_products.sum()))),
        (
            "直通车优品加推数不足客户数",
            f"{COL_CASH_DAY} > 0 且 {COL_PREMIUM_PRODUCTS} < {PREMIUM_PRODUCTS_MIN}",
            float(int(row_metrics["_is_premium_insufficient"].sum())),
        ),
        (
            "当日商机成本高于500元客户数",
            f"{COL_CASH_DAY} / {COL_INQUIRY} > {OPPORTUNITY_COST_THRESHOLD}（且询盘数>0）",
            float(int(row_metrics["_is_high_opportunity_cost"].sum())),
        ),
    ]


def compute_statistics(
    df_today: pd.DataFrame,
    df_yesterday: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """基于今日/昨日推广数据计算统计；有昨日数据时增加差值列。"""
    today_rows = _calc_metric_rows(df_today)
    names = [r[0] for r in today_rows]
    descs = [r[1] for r in today_rows]
    today_vals = [r[2] for r in today_rows]

    if df_yesterday is None:
        return pd.DataFrame({"指标": names, "说明": descs, "数值": today_vals})

    yesterday_map = {r[0]: r[2] for r in _calc_metric_rows(df_yesterday)}
    diffs = [round(t - yesterday_map.get(n, 0), 4) for n, t in zip(names, today_vals)]

    return pd.DataFrame(
        {
            "指标": names,
            "说明": descs,
            "今日数值": today_vals,
            "昨日数值": [yesterday_map.get(n, 0) for n in names],
            "差值(今日-昨日)": diffs,
        }
    )


def default_output_filename(
    table1_path: Path,
    original_filename: str | None = None,
) -> str:
    """默认输出：今日推广数据原始文件名（无扩展名）+ 广告数据.xlsx"""
    if original_filename:
        stem = Path(original_filename).stem
    else:
        stem = Path(table1_path).stem
    return f"{stem}广告数据.xlsx"


def build_detail_df(
    df1: pd.DataFrame,
    output_df: pd.DataFrame,
    row_metrics: pd.DataFrame,
    key1: str,
) -> pd.DataFrame:
    """组装用于战区/销售汇总的明细表。"""
    detail = output_df.copy()
    detail[COL_SALES] = df1[COL_SALES].astype(str).str.strip()
    detail[COL_CASH_DAY] = to_num(df1[COL_CASH_DAY])
    detail[COL_BUDGET] = to_num(df1[COL_BUDGET])
    detail["_is_active"] = row_metrics["_is_active"].values
    detail["_is_low_balance"] = row_metrics["_is_low_balance"].values
    detail["_is_high_opportunity_cost"] = row_metrics["_is_high_opportunity_cost"].values
    return detail


def aggregate_report_metrics(df: pd.DataFrame) -> dict:
    """计算汇报用五项指标。"""
    total_consume = df[COL_CASH_DAY].sum()
    total_budget = df[COL_BUDGET].sum()
    usage_rate = total_consume / total_budget if total_budget else 0
    return {
        "总消耗": round(total_consume, 2),
        "消耗使用率": usage_rate,
        "活跃客户数": int(df["_is_active"].sum()),
        "余额不足客户数": int(df["_is_low_balance"].sum()),
        "商机成本高于500客户数": int(df["_is_high_opportunity_cost"].sum()),
    }


def format_metric_value(name: str, value) -> str:
    if name == "消耗使用率":
        return f"{value:.2%}"
    if name == "总消耗":
        return f"{value:,.2f}"
    return str(value)


def is_valid_sales_name(name: str) -> bool:
    return bool(name) and name not in ("-", "nan")


def get_unassigned_sales_df(zone_df: pd.DataFrame) -> pd.DataFrame:
    """销售姓名为空、'-' 或 nan 的客户。"""
    sales = zone_df[COL_SALES].astype(str).str.strip()
    return zone_df[~sales.apply(is_valid_sales_name)]


def is_valid_ops_name(name: str) -> bool:
    return bool(name) and name not in ("-", "nan")


def get_unassigned_ops_df(detail: pd.DataFrame) -> pd.DataFrame:
    """运营为空或未映射到柯桥/嵊州的客户。"""
    ops_raw = detail[COL_OPS]
    ops = ops_raw.astype(str).str.strip()
    return detail[
        ops_raw.isna()
        | (~ops.apply(is_valid_ops_name))
        | (~ops.isin(OPS_TO_REGION))
    ]


def ops_names_in_region(region: str) -> list:
    return sorted(name for name, r in OPS_TO_REGION.items() if r == region)


def write_region_member_report(
    writer: pd.ExcelWriter,
    detail: pd.DataFrame,
    key1: str,
    col_company: str,
    *,
    sheet_name: str,
    region_order: list,
    region_col: str,
    member_col: str,
    member_prefix: str,
    get_region_unassigned=None,
    global_unassigned=None,
    global_unassigned_title: str = "【未分配】",
) -> None:
    """按区域 -> 成员写入卡片式汇报（销售战区 / 运营区域通用）。"""
    ws = writer.book.create_sheet(sheet_name)

    title_font = Font(bold=True, size=14)
    region_font = Font(bold=True, size=12, color="1F4E79")
    member_font = Font(bold=True, size=11, color="333333")
    header_font = Font(bold=True)
    region_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    member_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    detail_fill = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")

    row_idx = 1

    def write_row(cells, font=None, fill=None):
        nonlocal row_idx
        for col_idx, val in enumerate(cells, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row_idx += 1

    def write_metrics_block(metrics: dict, font=None, fill=None):
        for name in REPORT_METRICS:
            write_row([name, format_metric_value(name, metrics[name])], font=font, fill=fill)

    def write_low_balance_detail(sub: pd.DataFrame):
        write_row(["余额不足客户明细"], font=header_font, fill=detail_fill)
        write_row([key1, col_company, COL_BALANCE], font=header_font)
        for _, r in sub.iterrows():
            write_row([r[key1], r[col_company], r[COL_BALANCE]])

    def write_high_cost_detail(sub: pd.DataFrame):
        write_row(["商机成本高于500客户明细"], font=header_font, fill=detail_fill)
        write_row([key1, col_company], font=header_font)
        for _, r in sub.iterrows():
            write_row([r[key1], r[col_company]])

    def write_member_block(member_df: pd.DataFrame, title: str):
        if member_df.empty:
            return
        metrics = aggregate_report_metrics(member_df)
        write_row([title], font=member_font, fill=member_fill)
        write_metrics_block(metrics, font=member_font, fill=member_fill)
        if metrics["余额不足客户数"] > 0:
            write_low_balance_detail(member_df[member_df["_is_low_balance"]])
        if metrics["商机成本高于500客户数"] > 0:
            write_high_cost_detail(member_df[member_df["_is_high_opportunity_cost"]])
        write_row([])

    for region in region_order:
        region_df = detail[detail[region_col] == region]
        if region_df.empty:
            continue

        write_row([f"【{region}】"], font=title_font, fill=region_fill)
        write_metrics_block(aggregate_report_metrics(region_df), font=region_font, fill=region_fill)
        write_row([])

        if member_col == COL_OPS:
            member_names = ops_names_in_region(region)
        else:
            member_names = sorted(
                str(n).strip()
                for n in region_df[member_col].unique()
                if is_valid_sales_name(str(n).strip())
            )

        for name in member_names:
            member_df = region_df[region_df[member_col].astype(str).str.strip() == name]
            write_member_block(member_df, f"【{member_prefix}：{name}】")

        if get_region_unassigned:
            unassigned_df = get_region_unassigned(region_df)
            if not unassigned_df.empty:
                write_member_block(unassigned_df, f"【未分配{member_prefix}】")

        write_row([])

    if global_unassigned:
        unassigned_all = global_unassigned(detail)
        if not unassigned_all.empty:
            write_member_block(unassigned_all, global_unassigned_title)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18


def write_warzone_sales_report(
    writer: pd.ExcelWriter,
    detail: pd.DataFrame,
    key1: str,
    col_company: str,
) -> None:
    write_region_member_report(
        writer,
        detail,
        key1,
        col_company,
        sheet_name=SHEET_SALES_REPORT,
        region_order=WARZONE_ORDER,
        region_col="战区",
        member_col=COL_SALES,
        member_prefix="销售",
        get_region_unassigned=get_unassigned_sales_df,
    )


def write_ops_region_report(
    writer: pd.ExcelWriter,
    detail: pd.DataFrame,
    key1: str,
    col_company: str,
) -> None:
    write_region_member_report(
        writer,
        detail,
        key1,
        col_company,
        sheet_name=SHEET_OPS_REPORT,
        region_order=OPS_REGION_ORDER,
        region_col=COL_OPS_REGION,
        member_col=COL_OPS,
        member_prefix="运营",
        global_unassigned=get_unassigned_ops_df,
        global_unassigned_title="【未分配运营】",
    )


def run_merge(
    table1_path: Path,
    table2_path: Path,
    table3_path: Path,
    output_path: Path,
    table1_yesterday_path: Path | None = None,
    log=print,
) -> Path:
    """执行三表合并并写入 Excel，返回输出路径。"""
    table1_path = Path(table1_path)
    table2_path = Path(table2_path)
    table3_path = Path(table3_path)
    table1_yesterday_path = Path(table1_yesterday_path) if table1_yesterday_path else None
    output_path = Path(output_path)

    for p, label in [
        (table1_path, "表1"),
        (table2_path, "表2"),
        (table3_path, "表3"),
    ]:
        if not p.is_file():
            raise FileNotFoundError(f"{label} 文件不存在: {p}")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    key1 = "admin_mbr_id"
    key2 = "主账号ID"
    col_company = "公司名称"
    col_ops = "运营"
    col_supervisor_src = "主管姓名"  # 表1 原始列名
    col_supervisor_out = "主管名称"  # 输出列名
    col_warzone = "战区"

    stat_cols = [
        COL_CASH_DAY,
        COL_BUDGET,
        COL_FIN_DAY,
        COL_CASH_MONTH,
        COL_BALANCE,
        COL_PREMIUM_PRODUCTS,
        COL_SHOP_PRODUCTS,
        COL_SHOP_PREMIUM,
        COL_INQUIRY,
    ]

    df1 = pd.read_excel(table1_path, sheet_name=0)
    df1_yesterday = None
    if table1_yesterday_path:
        if not table1_yesterday_path.is_file():
            raise FileNotFoundError(f"昨日推广数据文件不存在: {table1_yesterday_path}")
        df1_yesterday = pd.read_excel(table1_yesterday_path, sheet_name=0)
        for col in stat_cols + [key1, col_supervisor_src, COL_SALES]:
            if col not in df1_yesterday.columns:
                raise KeyError(f"昨日推广数据缺少列: {col}")

    df2 = pd.read_excel(table2_path, sheet_name=0)
    xl3 = pd.ExcelFile(table3_path)
    sheet3 = "28号" if "28号" in xl3.sheet_names else xl3.sheet_names[0]
    df3 = pd.read_excel(table3_path, sheet_name=sheet3)

    for col, name in [
        (key1, "表1"),
        (col_supervisor_src, "表1"),
        (COL_SALES, "表1"),
        *[(c, "表1") for c in stat_cols],
        (key2, "表2"),
        (key1, "表3"),
    ]:
        if col not in (df1 if name == "表1" else df2 if name == "表2" else df3).columns:
            raise KeyError(f"{name} 缺少列: {col}")

    # 表2、表3 按主键去重（保留第一条）
    df2_u = df2[[key2, col_company]].drop_duplicates(subset=[key2], keep="first")
    df3_u = df3[[key1, col_ops]].drop_duplicates(subset=[key1], keep="first")

    df2_u = df2_u.rename(columns={key2: key1})
    df2_u["_key"] = normalize_id(df2_u[key1])
    df3_u["_key"] = normalize_id(df3_u[key1])

    # 以表1 为基准（保留表1 全部行）
    result = df1[[key1, col_supervisor_src]].copy()
    result = result.rename(columns={col_supervisor_src: col_supervisor_out})
    result["_key"] = normalize_id(result[key1])

    result = result.merge(
        df2_u[["_key", col_company]],
        on="_key",
        how="left",
    )
    result = result.merge(
        df3_u[["_key", col_ops]],
        on="_key",
        how="left",
    )
    result = result.drop(columns=["_key"])
    result[col_warzone] = assign_warzone(result[col_supervisor_out])
    result[COL_OPS_REGION] = assign_ops_region(result[col_ops])

    row_metrics = prepare_row_metrics(df1)
    result = pd.concat([result, row_metrics], axis=1)

    # 统计数据所用原字段（来自表1）
    cash_day = to_num(df1[COL_CASH_DAY])
    budget_day = to_num(df1[COL_BUDGET])
    result[COL_CASH_DAY] = cash_day.round(2)
    result[COL_BUDGET] = budget_day.round(2)
    result[COL_FIN_DAY] = to_num(df1[COL_FIN_DAY]).round(2)
    result[COL_CASH_MONTH] = pd.to_numeric(df1[COL_CASH_MONTH], errors="coerce").round(2)
    result[COL_INQUIRY] = to_num(df1[COL_INQUIRY]).round(2)
    result["消耗使用率"] = (cash_day / budget_day.replace(0, float("nan"))).round(4)

    result[COL_SHOP_PRODUCTS] = to_num(df1[COL_SHOP_PRODUCTS]).round(0)
    result[COL_SHOP_PREMIUM] = to_num(df1[COL_SHOP_PREMIUM]).round(0)

    base_cols = [key1, col_supervisor_out, col_warzone, col_company, col_ops, COL_OPS_REGION]
    shop_cols = [COL_SHOP_PRODUCTS, COL_SHOP_PREMIUM]
    raw_stat_cols = [
        COL_CASH_DAY,
        COL_BUDGET,
        COL_FIN_DAY,
        COL_CASH_MONTH,
        COL_INQUIRY,
        COL_BALANCE,
        COL_PREMIUM_PRODUCTS,
    ]
    derived_cols = [
        "消耗使用率",
        "当日卡券消耗",
        "余额可用天数",
        "当日商机成本",
        "备注",
    ]
    full_out_cols = base_cols + shop_cols + raw_stat_cols + derived_cols
    output_df = result[full_out_cols]
    stats = compute_statistics(df1, df1_yesterday)
    detail = build_detail_df(df1, output_df, row_metrics, key1)

    matched_company = output_df[col_company].notna().sum()
    matched_ops = output_df[col_ops].notna().sum()
    total = len(output_df)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        output_df.to_excel(writer, sheet_name=SHEET_BASE, index=False)
        stats.to_excel(writer, sheet_name=SHEET_STATS, index=False)
        write_warzone_sales_report(writer, detail, key1, col_company)
        write_ops_region_report(writer, detail, key1, col_company)
        set_column_widths(writer.sheets[SHEET_BASE], output_df)
        set_column_widths(writer.sheets[SHEET_STATS], stats, min_width=14, max_width=45)

    unmapped = output_df[output_df[col_warzone].isna()][col_supervisor_out].dropna().unique()

    log(f"今日推广数据行数: {len(df1)}")
    if df1_yesterday is not None:
        log(f"昨日推广数据行数: {len(df1_yesterday)}（已输出今日/昨日/差值统计）")
    log(f"输出行数: {total}")
    log(f"匹配到公司名称: {matched_company} ({matched_company/total:.1%})")
    log(f"匹配到运营: {matched_ops} ({matched_ops/total:.1%})")
    log("战区分布:")
    log(output_df[col_warzone].value_counts(dropna=False).to_string())
    if len(unmapped):
        log(f"未匹配战区的主管 ({len(unmapped)} 人): {', '.join(map(str, unmapped))}")
    log("\n统计数据:")
    log(stats.to_string(index=False))
    log(f"\n已生成工作表: {SHEET_BASE}、{SHEET_STATS}、{SHEET_SALES_REPORT}、{SHEET_OPS_REPORT}")
    log(f"\n已保存: {output_path}")
    return output_path


def main() -> Path:
    return run_merge(TABLE1_PATH, TABLE2_PATH, TABLE3_PATH, OUTPUT_PATH)


if __name__ == "__main__":
    main()
